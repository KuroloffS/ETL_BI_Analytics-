import os
import pandas as pd
import numpy as np
import psycopg2
from psycopg2 import sql, pool
from openpyxl import load_workbook
from dotenv import load_dotenv

# Загрузка переменных окружения
load_dotenv()

# Реквизиты подключения к базе данных
DB_NAME = os.getenv('DB_NAME')
DB_HOST = os.getenv('DB_HOST')
DB_USER = os.getenv('DB_USER')
DB_PASSWORD = os.getenv('DB_PASSWORD')

# Путь к Excel-файлу
#SOURCE_FILE = r'\\10.1.19.245\C2M Logistics\BI Analytics\Plan\Plan 2024 for dashboard v1.xlsx'
SOURCE_FILE = r'\\10.1.19.245\C2M Logistics\BI Analytics\tariff.xlsx'
db_target_table = 'tariff'

def truncate_table(cursor, table_name):
    """Удаляет все записи из таблицы."""
    try:
        cursor.execute(sql.SQL("TRUNCATE TABLE {table}").format(
            table=sql.Identifier(table_name)
        ))
        print(f"Таблица {table_name} успешно очищена.")
    except Exception as e:
        raise Exception(f"Ошибка при очистке таблицы {table_name}: {e}")

def insert_data(cursor, data, table_name):
    """Вставляет данные в таблицу."""
    try:
        # Получение имен столбцов и приведение их к нижнему регистру
        columns = [col.lower() for col in data.columns]

        # Подготовка запроса INSERT
        insert_query = sql.SQL("""
        INSERT INTO {table} ({columns})
        VALUES ({placeholders})
        """).format(
            table=sql.Identifier(table_name),
            columns=sql.SQL(', ').join(map(sql.Identifier, columns)),
            placeholders=sql.SQL(', ').join(sql.Placeholder() for _ in columns)
        )

        # Выполнение вставки данных
        cursor.executemany(insert_query, data.values.tolist())
        print(f"Данные успешно записаны в таблицу {table_name}.")
    except Exception as e:
        raise Exception(f"Ошибка при вставке данных в таблицу {table_name}: {e}")

def single_excel_to_db():
    try:
        # Создание пула подключений
        db_pool = psycopg2.pool.SimpleConnectionPool(
            1, 10,
            dbname=DB_NAME,
            host=DB_HOST,
            user=DB_USER,
            password=DB_PASSWORD
        )

        # Получение соединения из пула
        with db_pool.getconn() as conn:
            with conn.cursor() as cursor:
                # Чтение данных из Excel
                raw_df = pd.read_excel(SOURCE_FILE, sheet_name=0, engine='openpyxl')
                

                if raw_df.empty:
                    print(f"[ERROR] Файл Excel пуст: {len(raw_df)} строк.")
                    return
                raw_df.columns = raw_df.columns.str.lower()
                print(raw_df.head())
                
                # Применение fillna(0.00) ко всем числовым столбцам
                numeric_columns = raw_df.select_dtypes(include=['number']).columns
                raw_df[numeric_columns] = raw_df[numeric_columns].fillna(0.00)
                
                #raw_df['stage'] = ''

                raw_df = raw_df[['liter', 'weight', 'country', 'service', 'per_kg', 'per_item', 'datefrom','tariff_range', 'stage']] #, 
                raw_df['stage'] = raw_df['stage'].replace({np.nan: None})
            

                # Очистка таблицы перед записью
                truncate_table(cursor, db_target_table)

                # Запись данных в таблицу
                insert_data(cursor, raw_df, db_target_table)

                # Фиксация изменений
                conn.commit()

    except Exception as e:
        print(f"Ошибка выполнения: {e}")
    finally:
        # Закрытие пула подключений
        if 'db_pool' in locals() and db_pool:
            db_pool.closeall()

if __name__ == "__main__":
    single_excel_to_db()
