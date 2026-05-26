import os
import re
import psycopg2
from psycopg2 import sql
from datetime import datetime
from openpyxl import load_workbook
from dotenv import load_dotenv
import logging

# Настройка логирования
logging.basicConfig(
    filename="process_to_registry.log",
    level=logging.DEBUG,
    format="%(asctime)s - %(levelname)s - %(message)s",
    encoding="utf-8",  # Устанавливаем кодировку UTF-8
)

# Загрузка переменных окружения
load_dotenv()

# Реквизиты подключения к базе данных
DB_NAME = os.getenv('DB_NAME')
DB_HOST = os.getenv('DB_HOST')
DB_USER = os.getenv('DB_USER')
DB_PASSWORD = os.getenv('DB_PASSWORD')

# Начальная директория поиска
START_DIR = r'\\10.1.19.245\C2M Logistics\2025\CAINIAO'
TABLE_NAME = 'excel_registry'

# Шаблон для файлов
FILE_PATTERN = re.compile(r".*\.xlsx$")

# Регулярное выражение для разбора target_folder для поиска номера партии и даты
PACKAGE_PATTERN = re.compile(r"-\s*(\d+).*?\((\d{2}\.\d{2}\.\d{4})\)")

def check_required_columns(headers):
    required_columns = ["cartonid", "itemid", "masterno", "ordercode", "consigneecountry", "service", "danger_type"]
    missing_columns = [col for col in required_columns if col not in headers]
    weight_column_present = any("weight" in col for col in headers)
    
    if missing_columns or not weight_column_present:
        raise ValueError(f"Missing columns: {', '.join(missing_columns)} or 'weight' column not found.")
    return headers

def process_excel_file(file_path):
    try:
        logging.info(f"Processing file: {file_path}")
        workbook = load_workbook(filename=file_path, read_only=True)
        sheet = workbook.active

        if sheet.max_row < 1:
            raise ValueError(f"File {file_path} has no data")

        headers = [cell.value.lower() for cell in sheet[1] if cell.value]
        if not headers:
            raise ValueError(f"First row in file {file_path} is empty")

        check_required_columns(headers)
        logging.info(f"File {file_path} passed column validation")
        return headers
    except Exception as e:
        logging.error(f"Error processing file {file_path}: {e}")
        raise

def is_file_in_registry(cursor, file_path):
    query = f"SELECT EXISTS (SELECT 1 FROM {TABLE_NAME} WHERE full_path = %s);"
    cursor.execute(query, (file_path,))
    return cursor.fetchone()[0]

def update_registry_entry(cursor, conn, file_path, is_error, error_message):
    update_query = f"""
        UPDATE {TABLE_NAME}
        SET is_error = %s, error_message = %s
        WHERE full_path = %s
    """
    cursor.execute(update_query, (is_error, error_message, file_path))
    conn.commit()

def find_files_and_save_to_registry(start_dir):
    data_to_insert = []
    try:
        conn = psycopg2.connect(
            dbname=DB_NAME,
            host=DB_HOST,
            user=DB_USER,
            password=DB_PASSWORD
        )
        cursor = conn.cursor()

        for root, dirs, files in os.walk(start_dir):
            for file in files:
                if file.startswith('~$'):
                    continue

                if FILE_PATTERN.match(file):
                    file_path = os.path.join(root, file)
                    file_stat = os.stat(file_path)
                    file_dt_create = datetime.fromtimestamp(file_stat.st_ctime).strftime('%Y-%m-%d %H:%M:%S')
                    logging.info(f"Found file: {file_path}")

                    if not is_file_in_registry(cursor, file_path):
                        target_folder = os.path.basename(root)
                        file_name = os.path.basename(file_path)

                        try:
                            headers = process_excel_file(file_path)
                            file_size_kb = file_stat.st_size // 1024
                            dtflow = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
                            field_count = len(headers)
                            field_list = ";".join(headers)
                            is_new = True
                            is_error = False
                            error_message = None
                            dt_package, package_num = None, None

                            match = PACKAGE_PATTERN.search(file_path)
                            if match:
                                package_num = match.group(1)
                                dt_package = datetime.strptime(match.group(2), '%d.%m.%Y').strftime('%Y-%m-%d')

                            data_to_insert.append(
                                (target_folder, file_name, file_path, file_dt_create, file_size_kb, dtflow, field_count, field_list, is_new, is_error, error_message, dt_package, package_num)
                            )
                        except Exception as e:
                            logging.error(f"Error processing file {file_path}: {e}")
                            is_new = False
                            is_error = True
                            error_message = str(e)
                            dt_package, package_num = None, None

                            data_to_insert.append(
                                (target_folder, file_name, file_path, None, None, datetime.now().strftime('%Y-%m-%d %H:%M:%S'), 0, None, is_new, is_error, error_message, dt_package, package_num)
                            )

        if data_to_insert:
            save_to_db(cursor, conn, data_to_insert)
        else:
            logging.info("No new files found.")

    except Exception as e:
        logging.error(f"Error during file processing: {e}")
    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()

def save_to_db(cursor, conn, data):
    try:
        insert_query = f"""
            INSERT INTO {TABLE_NAME} (
                target_folder, file_name, full_path, dt_create, size_kb, dtflow, field_count, field_list, is_new, is_error, error_message, dt_package, package_num
            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
        """
        cursor.executemany(insert_query, data)
        conn.commit()
        logging.info(f"Successfully inserted {len(data)} rows into {TABLE_NAME}")
    except Exception as e:
        logging.error(f"Error saving data to database: {e}")
        conn.rollback()

if __name__ == "__main__":
    logging.info("Starting process...")
    find_files_and_save_to_registry(START_DIR)
    logging.info("Process completed.")
